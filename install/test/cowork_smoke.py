"""Generate, reopen and render real office/PDF/image files in the release image."""

from pathlib import Path
import subprocess
import sys

import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas


def run(*args):
    return subprocess.run(args, check=True, text=True, capture_output=True, timeout=60).stdout


def main(root):
    font = "Noto Sans CJK SC"
    assert "NotoSansCJK" in run("fc-match", "-f", "%{file}", font)
    frame = pd.DataFrame({"Item": ["Alpha", "Beta"], "Value": [20, 22]})
    assert frame.Value.sum() == 42
    frame.plot.bar(x="Item", y="Value", legend=False, color="#235b82")
    plt.title("Cowork chart")
    plt.tight_layout()
    plt.savefig(root / "chart.png")
    plt.close()
    with Image.open(root / "chart.png") as image:
        assert image.width >= 300 and image.height >= 300

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    document.add_heading("Cowork document", level=1)
    document.add_paragraph("Document marker. 中文文档与图表。")
    document.add_picture(str(root / "chart.png"), width=Inches(5))
    document.save(root / "document.docx")
    assert "Document marker" in Document(root / "document.docx").paragraphs[1].text

    frame.to_excel(root / "workbook.xlsx", index=False, engine="openpyxl")
    workbook = load_workbook(root / "workbook.xlsx")
    sheet = workbook.active
    sheet["A4"] = "Spreadsheet marker"
    sheet["B4"] = "=SUM(B2:B3)"
    sheet.column_dimensions["A"].width = 28
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    workbook.save(root / "workbook.xlsx")
    assert load_workbook(root / "workbook.xlsx").active["B4"].value == "=SUM(B2:B3)"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Presentation marker"
    slide.shapes.add_picture(str(root / "chart.png"), Inches(2), Inches(1.5), width=Inches(6))
    presentation.save(root / "presentation.pptx")
    assert Presentation(root / "presentation.pptx").slides[0].shapes.title.text == "Presentation marker"

    pdf = canvas.Canvas(str(root / "generated.pdf"))
    pdf.drawString(72, 760, "PDF marker")
    pdf.drawImage(str(root / "chart.png"), 72, 360, width=400, height=300)
    pdf.save()

    expected = {"document": "Document marker", "workbook": "Spreadsheet marker", "presentation": "Presentation marker"}
    for name, extension in (("document", "docx"), ("workbook", "xlsx"), ("presentation", "pptx")):
        run("libreoffice", f"-env:UserInstallation={(root / 'office-profile').as_uri()}",
            "--headless", "--convert-to", "pdf", "--outdir", str(root), str(root / f"{name}.{extension}"))
    expected["generated"] = "PDF marker"
    for name, marker in expected.items():
        path = root / f"{name}.pdf"
        reader = PdfReader(path)
        assert len(reader.pages) == 1, (name, len(reader.pages))
        assert marker in reader.pages[0].extract_text(), name
        text = run("pdftotext", str(path), "-")
        assert marker in text, name
        if name == "document":
            assert "中文文档与图表" in text, "CJK text must survive Office rendering"
        if name == "workbook":
            assert "42" in text, "LibreOffice must recalculate the workbook formula"
        run("pdftoppm", "-singlefile", "-scale-to", "1200", "-png", str(path), str(root / f"{name}-render"))
        with Image.open(root / f"{name}-render.png") as image:
            assert max(image.size) == 1200
            assert image.convert("L").getextrema()[0] < 100, "rendered page must contain visible content"
    print("DOCX/XLSX/PPTX generated, reopened and rendered; formulas, PDF text, chart and CJK font verified")


if __name__ == "__main__":
    main(Path(sys.argv[1]).resolve())
